# extract.py
# Stephan DeLuna and Jocelyn Rogers
# extracts text from documents

# - - - - - - - - - - #

import os
from pypdf import PdfReader
from docx import Document

# - - - - - - - - - - #

# load .txt
def loadTxt(path):
	encodings = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
	for encodingName in encodings:
		try:
			with open(path, "r", encoding=encodingName) as f:
				return f.read()
		except UnicodeDecodeError:
			continue
	raise UnicodeDecodeError("Could not decode text file with supported encodings")

# load .pdf
def loadPdf(path):
	reader = PdfReader(path)
	text = ""

	for page in reader.pages:
		extracted = page.extract_text()
		if extracted:
			text += extracted + "\n"

	return text

# load .docx
def loadDocx(path):
	doc = Document(path)
	paragraphs = [p.text for p in doc.paragraphs]
	return "\n".join(paragraphs)

# - - - - - - - - - - #

# unified loader
def loadDocument(path):
	ext = os.path.splitext(path)[1].lower()

	if ext == ".txt":
		return loadTxt(path)
	elif ext == ".pdf":
		return loadPdf(path)
	elif ext == ".docx":
		return loadDocx(path)
	else:
		raise ValueError(f"Unsupported file type: {ext}")

# - - - - - - - - - - #


