# output.py
# Stephan DeLuna
# outputs final summaries as survey paper files

# - - - - - - - - - - #

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

# - - - - - - - - - - #

# build paper body from dictionary
def buildBodyText(summaryDict):
	bodyText = ""
	for sectionName, sectionText in summaryDict.items():
		if sectionName == "paper":
			bodyText += "Survey Summary\n"
		else:
			bodyText += sectionName.title() + "\n"
		bodyText += sectionText.strip() + "\n\n"
	return bodyText.strip()

# build the survey paper text
def buildSurveyPaper(
	title,
	authorName,
	introductionText,
	summaryDict,
	conclusionText,
	abstractText=""
):
	paperText = ""
	# title and author
	paperText += title.strip() + "\n"
	paperText += authorName.strip() + "\n\n"
	# abstract only if provided
	if abstractText.strip() != "":
		paperText += "Abstract\n"
		paperText += abstractText.strip() + "\n\n"
	# introduction
	paperText += "Introduction\n"
	paperText += introductionText.strip() + "\n\n"
	# body
	bodyText = buildBodyText(summaryDict)
	if bodyText != "":
		paperText += bodyText + "\n\n"
	# conclusion
	paperText += "Conclusion\n"
	paperText += conclusionText.strip() + "\n"
	return paperText

# save as txt
def saveTxt(path, paperText):
	with open(path, "w", encoding="utf-8") as f:
		f.write(paperText)

# save as docx
def saveDocx(
	path,
	title,
	authorName,
	introductionText,
	summaryDict,
	conclusionText,
	abstractText=""
):
	doc = Document()
	# title and author
	doc.add_heading(title.strip(), 0)
	doc.add_paragraph(authorName.strip())
	# abstract only if provided
	if abstractText.strip() != "":
		doc.add_heading("Abstract", level=1)
		doc.add_paragraph(abstractText.strip())
	# introduction
	doc.add_heading("Introduction", level=1)
	doc.add_paragraph(introductionText.strip())
	# body
	for sectionName, sectionText in summaryDict.items():
		if sectionName == "paper":
			doc.add_heading("Survey Summary", level=1)
		else:
			doc.add_heading(sectionName.title(), level=1)
		doc.add_paragraph(sectionText.strip())
	# conclusion
	doc.add_heading("Conclusion", level=1)
	doc.add_paragraph(conclusionText.strip())
	doc.save(path)

# save as pdf
def savePdf(
	path,
	title,
	authorName,
	introductionText,
	summaryDict,
	conclusionText,
	abstractText=""
):
	pdf = canvas.Canvas(path, pagesize=letter)
	pageWidth, pageHeight = letter
	x = 72
	y = pageHeight - 72
	lineHeight = 14
	def writeLine(text, fontName="Times-Roman", fontSize=12, extraGap=0):
		nonlocal y
		pdf.setFont(fontName, fontSize)
		words = text.split()
		currentLine = ""
		for word in words:
			testLine = currentLine + word + " "
			if pdf.stringWidth(testLine, fontName, fontSize) < (pageWidth - 144):
				currentLine = testLine
			else:
				pdf.drawString(x, y, currentLine.strip())
				y -= lineHeight
				currentLine = word + " "
				if y < 72:
					pdf.showPage()
					y = pageHeight - 72
					pdf.setFont(fontName, fontSize)
		if currentLine.strip() != "":
			pdf.drawString(x, y, currentLine.strip())
			y -= lineHeight
			if y < 72:
				pdf.showPage()
				y = pageHeight - 72
		y -= extraGap
	# title and author
	writeLine(title.strip(), "Times-Bold", 16, extraGap=4)
	writeLine(authorName.strip(), "Times-Roman", 12, extraGap=8)
	# abstract only if provided
	if abstractText.strip() != "":
		writeLine("Abstract", "Times-Bold", 14, extraGap=2)
		writeLine(abstractText.strip(), "Times-Roman", 12, extraGap=8)
	# introduction
	writeLine("Introduction", "Times-Bold", 14, extraGap=2)
	writeLine(introductionText.strip(), "Times-Roman", 12, extraGap=8)
	# body
	for sectionName, sectionText in summaryDict.items():
		if sectionName == "paper":
			writeLine("Survey Summary", "Times-Bold", 14, extraGap=2)
		else:
			writeLine(sectionName.title(), "Times-Bold", 14, extraGap=2)
		writeLine(sectionText.strip(), "Times-Roman", 12, extraGap=8)
	# conclusion
	writeLine("Conclusion", "Times-Bold", 14, extraGap=2)
	writeLine(conclusionText.strip(), "Times-Roman", 12, extraGap=0)
	pdf.save()

# export survey paper in chosen format
def exportSurveyPaper(
	path,
	fileType,
	title,
	authorName,
	introductionText,
	summaryDict,
	conclusionText,
	abstractText=""
):
	paperText = buildSurveyPaper(
		title,
		authorName,
		introductionText,
		summaryDict,
		conclusionText,
		abstractText
	)
	if fileType == "txt":
		saveTxt(path, paperText)
	elif fileType == "docx":
		saveDocx(
			path,
			title,
			authorName,
			introductionText,
			summaryDict,
			conclusionText,
			abstractText
		)
	elif fileType == "pdf":
		savePdf(
			path,
			title,
			authorName,
			introductionText,
			summaryDict,
			conclusionText,
			abstractText
		)
	else:
		raise ValueError("Unsupported file type")

# - - - - - - - - - - #
